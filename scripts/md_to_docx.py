import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# Minimal Markdown to DOCX converter focusing on headings, paragraphs, lists, and tables
# This is not a full Markdown parser; tailored for our document structure.

def add_heading(doc, text, level):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if level == 1:
        run.bold = True
        run.font.size = Pt(16)
    elif level == 2:
        run.bold = True
        run.font.size = Pt(14)
    else:
        run.bold = True
        run.font.size = Pt(12)
    return p


def add_paragraph(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_list(doc, text):
    # simple bullet list
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(text)


def add_table_from_pipe(doc, header, rows):
    cols = len(header)
    table = doc.add_table(rows=1, cols=cols)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(header):
        hdr_cells[i].text = h.strip()
    for row in rows:
        cells = table.add_row().cells
        for i, cell in enumerate(row):
            cells[i].text = cell.strip()
    doc.add_paragraph('')


def parse_markdown(md_text):
    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    style.font.size = Pt(12)

    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith('## '):
            add_heading(doc, line[3:].strip(), 1)
        elif line.startswith('### '):
            add_heading(doc, line[4:].strip(), 2)
        elif line.startswith('#### '):
            add_heading(doc, line[5:].strip(), 3)
        elif line.startswith('- '):
            add_list(doc, line[2:].strip())
        elif line.startswith('|') and '---' in line:
            # header separator line for table will be handled when seen with context
            i += 1
            continue
        elif line.startswith('|'):
            # Parse GitHub-style pipe tables (header + separator + rows)
            # Find header
            header_line = line
            # Next line should be separator
            sep_line = lines[i+1] if i+1 < len(lines) else ''
            # Accumulate rows until a blank or non-pipe line
            row_start = i + 2
            row_lines = []
            j = row_start
            while j < len(lines) and lines[j].startswith('|'):
                row_lines.append(lines[j])
                j += 1
            header = [c.strip() for c in header_line.strip('|').split('|')]
            rows = [[c.strip() for c in rl.strip('|').split('|')] for rl in row_lines]
            add_table_from_pipe(doc, header, rows)
            i = j
            continue
        else:
            if line.strip() == '':
                doc.add_paragraph('')
            else:
                add_paragraph(doc, line.strip())
        i += 1

    return doc


def main():
    if len(sys.argv) != 3:
        print('Usage: python md_to_docx.py <input.md> <output.docx>')
        sys.exit(1)
    in_path = Path(sys.argv[1])
    out_path = Path(sys.argv[2])
    md = in_path.read_text(encoding='utf-8')
    doc = parse_markdown(md)
    doc.save(str(out_path))
    print(f'Wrote {out_path}')


if __name__ == '__main__':
    main()
